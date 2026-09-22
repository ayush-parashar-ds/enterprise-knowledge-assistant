from pathlib import Path
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from docx import Document as DocxDocument
from pptx import Presentation


def load_documents(folder_path):
    """
    Load PDF, DOCX, PPTX and TXT files from a folder.
    """

    documents = []
    folder = Path(folder_path)

    for file_path in folder.iterdir():

        # PDF files
        if file_path.suffix.lower() == ".pdf":
            loader = PyPDFLoader(str(file_path))
            docs = loader.load()

            for doc in docs:
                doc.metadata["source"] = file_path.name

            documents.extend(docs)

        # DOCX files
        elif file_path.suffix.lower() == ".docx":
            docx = DocxDocument(str(file_path))

            text = "\n".join(
                paragraph.text
                for paragraph in docx.paragraphs
                if paragraph.text.strip()
            )

            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": file_path.name}
                    )
                )

        # PPTX files
        elif file_path.suffix.lower() == ".pptx":
            presentation = Presentation(str(file_path))

            text_parts = []

            for slide in presentation.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)

            text = "\n".join(text_parts)

            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": file_path.name}
                    )
                )

        # TXT files
        elif file_path.suffix.lower() == ".txt":
            text = file_path.read_text(
                encoding="utf-8",
                errors="ignore"
            )

            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": file_path.name}
                    )
                )

    return documents